#!/usr/bin/env python3
"""Dựng slide (.pptx, lời giảng trong Notes) + bài giảng chi tiết (.md) cho MỘT buổi học từ một file nội dung.

Dùng: python3 build_session.py <noi_dung_buoi.py> <thu_muc_ra> [TIỀN_TỐ, mặc định W01]
File nội dung định nghĩa META (dict) và SLIDES (list) theo templates/session-content-schema.md.
Giao diện: META["theme"] (xem THEME_DEFAULT) — màu nền, chữ, các màu nhấn, phông. Không có thì dùng mặc định
trung tính (nền trắng, chữ xanh đậm). Màu chữ trên ô màu được chọn tự động theo độ tương phản WCAG.
Thời lượng slide chỉ giảng được TÍNH từ lời giảng (130 âm tiết/phút + 0,3 phút mỗi chỉ dẫn dừng/hỏi lớp).
Sau khi dựng, chạy scripts/qa_deck.py. Phụ thuộc: python-pptx.
"""
import importlib.util, math, re, sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

_spec = importlib.util.spec_from_file_location("session_content", sys.argv[1])
_mod = importlib.util.module_from_spec(_spec); _spec.loader.exec_module(_mod)
META, SLIDES = _mod.META, _mod.SLIDES
OUT = sys.argv[2]
PREFIX = sys.argv[3] if len(sys.argv) > 3 else "W01"
RATE = 130  # âm tiết/phút khi giảng bằng tiếng Việt (có ngắt nghỉ)
for _s in SLIDES:  # thời lượng slide chỉ giảng = tính từ lời giảng, không ước bằng tay
    if not _s.get("interact") and _s["kind"] not in ("break", "appendix", "refs", "title"):
        _spoken = re.sub(r"_\(.*?\)_", "", _s["notes"])
        _pause = 0.3 * len(re.findall(r"_\((dừng|hỏi lớp|chờ)", _s["notes"]))
        _s["min"] = max(1, math.ceil(len(re.findall(r"\w+", _spoken)) / RATE + _pause))

# ---------------- Giao diện ----------------
THEME_DEFAULT = dict(bg="FFFFFF", text="2D2D2D", title="1F4E79", muted="6B6B6B", panel="EAF1F8",
                     accents=["2E75B6", "D55E00", "009E73", "CC79A7", "0072B2", "E69F00"], font="Arial", dark_struct="1F4E79")
T = {**THEME_DEFAULT, **META.get("theme", {})}
def C(h): return RGBColor.from_string(h.lstrip("#").upper())
def lum(h):
    h = h.lstrip("#"); v = [int(h[i:i + 2], 16) / 255 for i in (0, 2, 4)]
    v = [x / 12.92 if x <= 0.03928 else ((x + 0.055) / 1.055) ** 2.4 for x in v]
    return 0.2126 * v[0] + 0.7152 * v[1] + 0.0722 * v[2]
def contrast(a, b):
    la, lb = sorted((lum(a), lum(b)), reverse=True); return (la + 0.05) / (lb + 0.05)
def on(fill):  # màu chữ đọc được trên một nền
    return max((T["text"], T["bg"], "FFFFFF", "1A1A1A"), key=lambda c: contrast(c, fill))
DARK = lum(T["bg"]) < 0.2
STRUCT_BG = T["bg"] if DARK else T["dark_struct"]   # nền của slide bìa/chia phần/kết luận
TEXT_ACC = [a for a in T["accents"] if contrast(a, T["bg"]) >= 4.5] or [T["title"]]  # màu nhấn dùng được cho chữ
_parts = list(dict.fromkeys(s["part"] for s in SLIDES))
def part_acc(part, text=False):
    a = T["accents"][_parts.index(part) % len(T["accents"])] if part in _parts else T["accents"][0]
    # Chữ dùng cùng màu với dải trang trí; màu nào không đủ tương phản (<4.5) thì chữ dùng màu tiêu đề.
    return a if not text or contrast(a, T["bg"]) >= 4.5 else T["title"]
FONT = T["font"]
BODY, TITLE, MUTED, LIGHT = C(T["text"]), C(T["title"]), C(T["muted"]), C(T["panel"])
NAVY = TITLE; BLUE = C(TEXT_ACC[0]); ACC = C(TEXT_ACC[1 % len(TEXT_ACC)])  # tên cũ, để file nội dung draw() dùng được

def plain(md):
    s = re.sub(r"\*\*(.+?)\*\*", r"\1", md)
    return re.sub(r"(?<!\w)_(.+?)_(?!\w)", r"\1", s)

def tb(slide, x, y, w, h, lines, size=24, color=None, bold=False, align=PP_ALIGN.LEFT, bullets=False, anchor=MSO_ANCHOR.TOP, bullet_color=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(12)
        runs = ([("•  ", bullet_color or color or BODY)] if bullets else []) + [(line, color or BODY)]
        for txt, col in runs:
            r = p.add_run(); r.text = txt
            r.font.size = Pt(size); r.font.name = FONT; r.font.color.rgb = col; r.font.bold = bold or (txt == "•  ")
    return box

def title(slide, text, size=28, color=None):
    tb(slide, 0.7, 0.3, 12.1, 1.2, [text], size=size, color=color or TITLE, bold=True, anchor=MSO_ANCHOR.MIDDLE)

def cite(slide, text):
    tb(slide, 0.7, 6.95, 11.4, 0.45, [text], size=12, color=MUTED)

def shape(slide, kind, x, y, w, h, fill=None, line=None, text=None, size=20, tcolor=None, bold=False):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is not None: s.fill.solid(); s.fill.fore_color.rgb = fill
    else: s.fill.background()
    if line is None and fill is not None: s.line.fill.background()
    else: s.line.color.rgb = line if line is not None else MUTED; s.line.width = Pt(2)
    s.shadow.inherit = False
    if text is not None:
        tf = s.text_frame; tf.word_wrap = True
        tf.paragraphs[0].text = text; tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        tc = tcolor if tcolor is not None else (C(on(str(fill))) if fill is not None else BODY)
        for r in tf.paragraphs[0].runs:
            r.font.size = Pt(size); r.font.name = FONT; r.font.color.rgb = tc; r.font.bold = bold
    return s

def arrow(slide, x1, y1, x2, y2, color=None):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color or BODY; c.line.width = Pt(2.5)
    ln = c.line._get_or_add_ln(); from pptx.oxml.ns import qn
    ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle"}))

def paint_bg(sl, hexcolor):
    sl.background.fill.solid(); sl.background.fill.fore_color.rgb = C(hexcolor)

BODY_CUR = []; CUR = {}
def visual(slide, v):
    acc = [C(a) for a in T["accents"]]
    if v.startswith("bars:"):
        items = [x.split("|") for x in v[5:].split(";")]; top = max(int(val) for _, val in items)
        y = 2.0
        for j, (lab, val) in enumerate(items):
            w = 9.0 * int(val) / top; col = acc[j % len(acc)]
            shape(slide, MSO_SHAPE.RECTANGLE, 2.6, y, w, 0.8, fill=col)
            tb(slide, 0.7, y + 0.1, 1.9, 0.6, [lab], size=22, bold=True, color=TITLE)
            tb(slide, 2.7 + w, y + 0.1, 1.5, 0.6, [val + "%"], size=24, bold=True, color=TITLE)
            y += 1.1
    elif v == "chips":
        heads = CUR.get("heads") or [str(j + 1) for j in range(len(BODY_CUR))]
        n = len(BODY_CUR); wbox = (12.1 - 0.2 * (n - 1)) / n
        for j, (b, c) in enumerate(zip(BODY_CUR, heads)):
            x = 0.7 + j * (wbox + 0.2)
            shape(slide, MSO_SHAPE.RECTANGLE, x, 1.8, wbox, 0.6, fill=acc[j % len(acc)], text=c, size=20, bold=True)
            shape(slide, MSO_SHAPE.RECTANGLE, x, 2.4, wbox, 2.6, fill=LIGHT, text=b, size=20)
    elif v.startswith("steps:"):
        labs = v[6:].split("|"); n = len(labs); wbox = (12.1 - 0.7 * (n - 1)) / n; x = 0.7
        for i, lab in enumerate(labs):
            shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, 2.6, wbox, 1.6, fill=acc[i % len(acc)], text=f"{i+1}. {lab}", size=24, bold=True)
            if i < n - 1: arrow(slide, x + wbox + 0.05, 3.4, x + wbox + 0.65, 3.4)
            x += wbox + 0.7
    elif v.startswith("twocol:"):
        a, b = v[7:].split("||")
        for j, (x, col) in enumerate(((0.7, a), (6.9, b))):
            head, items = col.split("|"); items = items.split(";")
            shape(slide, MSO_SHAPE.RECTANGLE, x, 1.7, 5.7, 0.7, fill=acc[j % len(acc)], text=head, size=22, bold=True)
            tb(slide, x + 0.1, 2.55, 5.5, 4.0, items, size=22, bullets=True, bullet_color=acc[j % len(acc)])
    elif v.startswith("grid:"):
        cells = [c.split("|") for c in v[5:].split(";")]  # "Tiêu đề|mô tả;…" — ô đầu được nhấn
        for i, (h, d) in enumerate(cells[:4]):
            x = 0.7 + (i % 2) * 6.15; y = 1.7 + (i // 2) * 1.75
            shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, 5.95, 1.55, fill=acc[0] if i == 0 else LIGHT, text=f"{h}\n{d}", size=20, bold=(i == 0))
    elif hasattr(_mod, "draw"):
        # Hình riêng của môn học: file nội dung tự vẽ bằng các hàm dựng sẵn.
        _mod.draw(slide, v, dict(shape=shape, tb=tb, arrow=arrow, cite=cite, MSO_SHAPE=MSO_SHAPE, on=lambda f: C(on(str(f))),
                                 NAVY=NAVY, BLUE=BLUE, BODY=BODY, MUTED=MUTED, LIGHT=LIGHT, ACC=ACC, RGBColor=RGBColor,
                                 TITLE=TITLE, ACCENTS=acc, TEXT_ACCENTS=[C(a) for a in TEXT_ACC]))
    else:
        raise ValueError(f"visual '{v}' không có trong builder chung và file nội dung không định nghĩa draw()")
    if CUR.get("caption"):
        tb(slide, 0.7, 5.35, 12.1, 0.8, [CUR["caption"]], size=22, color=C(TEXT_ACC[0]), bold=True)

def deco_strip(sl, y=7.3, h=0.2):  # dải màu nhấn trang trí
    w = 13.333 / len(T["accents"])
    for j, a in enumerate(T["accents"]):
        shape(sl, MSO_SHAPE.RECTANGLE, j * w, y, w if j == len(T["accents"]) - 1 else w + 0.01, h, fill=C(a))

def build_pptx():
    prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    n_ideas = len(META["ideas"])
    for i, s in enumerate(SLIDES, 1):
        sl = prs.slides.add_slide(blank); k = s["kind"]
        sl._element.cSld.set("name", k)  # qa_deck.py đọc loại slide từ đây
        pa, pat = C(part_acc(s["part"])), C(part_acc(s["part"], text=True))
        if k in ("title", "section", "conclusion", "break"):
            paint_bg(sl, STRUCT_BG); W = C(on(STRUCT_BG)); SOFT = C(T["muted"]) if DARK else C("B4C8DE")
            if k == "section":
                shape(sl, MSO_SHAPE.RECTANGLE, 0, 0, 0.35, 7.5, fill=pa)
                tb(sl, 0.9, 0.9, 3, 1.4, [f"0{s['idx']+1}"], size=72, color=pat if contrast(str(pat), STRUCT_BG) >= 3 else W, bold=True)
                tb(sl, 0.9, 2.4, 11.6, 2.0, [s["title"]], size=36, color=W, bold=True)
                for j, idea in enumerate(META["ideas"]):
                    tb(sl, 0.9, 4.6 + j * 0.6, 11.6, 0.6, [f"{j+1}. {idea}"], size=20,
                       color=W if j == s["idx"] else SOFT, bold=(j == s["idx"]))
            elif k == "title":
                tb(sl, 0.9, 1.8, 11.6, 1.4, [s["title"]], size=48, color=W, bold=True)
                shape(sl, MSO_SHAPE.RECTANGLE, 0.9, 3.35, 2.2, 0.12, fill=C(T["accents"][0]))
                tb(sl, 0.9, 3.7, 11.6, 2.5, s["body"], size=24, color=W)
                deco_strip(sl)
            elif k == "break":
                tb(sl, 0.8, 2.4, 11.7, 1.2, [s["title"]], size=44, color=W, bold=True, align=PP_ALIGN.CENTER)
                tb(sl, 0.8, 3.8, 11.7, 1.6, s["body"], size=24, color=W, align=PP_ALIGN.CENTER)
                deco_strip(sl)
            else:  # conclusion
                shape(sl, MSO_SHAPE.RECTANGLE, 0, 0, 0.35, 7.5, fill=pa)
                tb(sl, 0.9, 0.3, 12, 0.5, ["Kết luận"], size=20, color=pat if contrast(str(pat), STRUCT_BG) >= 4.5 else W, bold=True)
                tb(sl, 0.9, 0.85, 11.9, 1.3, [s["title"]], size=34, color=W, bold=True)
                tb(sl, 0.9, 2.4, 11.9, 3.6, s["body"], size=24, color=W, bullets=True, bullet_color=pat if contrast(str(pat), STRUCT_BG) >= 3 else W)
                if s.get("footer"): tb(sl, 0.9, 6.4, 11.9, 0.6, [s["footer"]], size=16, color=SOFT)
        else:
            if DARK: paint_bg(sl, T["bg"])
            shape(sl, MSO_SHAPE.RECTANGLE, 0, 0, 0.25 if k == "activity" else 0.12, 7.5, fill=pa)  # dải màu của phần
            title(sl, s["title"], size=26 if len(s["title"]) > 70 else 28, color=pat if k in ("activity", "hook") else TITLE)
            if k == "hook":
                for j, b in enumerate(s["body"]):
                    shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, 0.9 + j * 3.0, 2.6, 2.7, 2.0, fill=C(T["accents"][j % len(T["accents"])]), text=b, size=40, bold=True)
                if s.get("caption"): tb(sl, 0.9, 5.1, 11.5, 0.7, [s["caption"]], size=22, color=pat, bold=True)
            elif s.get("visual"):
                BODY_CUR[:] = s["body"]; CUR.clear(); CUR.update(s); visual(sl, s["visual"])
            elif k == "refs":
                tb(sl, 0.9, 1.75, 11.9, 5.0, s["body"], size=20)
            elif s["body"]:
                tb(sl, 0.9, 1.75, 11.9, 5.0, s["body"], size=24, bullets=(k not in ("outline",)), bullet_color=pat)
            if s.get("src") and not s.get("own_cite"):
                cite(sl, "Nguồn: " + re.sub(r"\s*·?\s*\[VERIFY[^\]]*\]", "", s["src"]))
            tb(sl, 12.2, 7.0, 1.0, 0.4, [str(i)], size=12, color=MUTED, align=PP_ALIGN.RIGHT)
        if s.get("alt"):
            cand = [sh for sh in sl.shapes if sh.shape_type == 1 and sh.width < prs.slide_width * 0.9 and sh.height < prs.slide_height * 0.9]
            tgt = cand[0] if cand else sl.shapes[-1]
            tgt._element.xpath("./*[1]/p:cNvPr")[0].set("descr", s["alt"])
        n = plain(s["notes"])
        if s.get("example"): n += "\n\nVÍ DỤ MINH HỌA: " + plain(s["example"])
        if s.get("interact"):
            it = s["interact"]
            n += f"\n\nTƯƠNG TÁC ({it['form']} · {it['time']}): {plain(it['task'])}\nĐáp án/kỳ vọng: {plain(it['answer'])}\nLời dẫn ra: {plain(it['out'])}"
        sl.notes_slide.notes_text_frame.text = f"[{s['part']} · {s['min']} phút]\n" + n
    prs.save(f"{OUT}/{PREFIX}_slides.pptx")

def build_md():
    total = sum(s["min"] for s in SLIDES)
    talk = sum(s["min"] for s in SLIDES if not s.get("interact") and s["kind"] not in ("break", "appendix"))
    act = sum(s["min"] for s in SLIDES if s.get("interact"))
    L = [f"# Bài giảng chi tiết – {META['session']}: {META['topic']}", "",
         f"**Học phần:** {META['code']} {META['course']} · **Thời lượng:** {META['minutes']} phút {META.get('minutes_note', '')} · **CLO:** {META['clo']} · **Giáo trình:** {META['gt']}",
         "", f"**Thông điệp cốt lõi:** {META['core']}", "",
         f"**{len(META['ideas'])} ý chính:** " + " · ".join(f"({i+1}) {x}" for i, x in enumerate(META["ideas"])), "",
         "**Sau buổi này sinh viên có thể:** " + "; ".join(META["outcomes"]), "",
         "## Tổng quan nhịp buổi học", "", "| Phút | Phần | Slide | Nội dung | Tương tác |", "|-------|------------|------|--------------------------|--------------|"]
    t = 0
    for i, s in enumerate(SLIDES, 1):
        if s["kind"] == "appendix": continue
        L.append(f"| {t}–{t+s['min']} | {s['part']} | {i} | {plain(s['title'])} | {s['interact']['form'] if s.get('interact') else '—'} |")
        t += s["min"]
    idx = {s["title"]: i for i, s in enumerate(SLIDES, 1)}
    def fill(t):  # {slide:Tiền tố tiêu đề} → số slide
        return re.sub(r"\{slide:(.+?)\}", lambda m: str(next(i for tt, i in idx.items() if tt.startswith(m.group(1)))), t)
    L += ["", f"Tổng kế hoạch: **{total} phút** / {META['minutes']} (dự phòng {META['minutes']-total} phút). Có tương tác: {act} phút; slide chỉ giảng: {talk} phút.",
          "", "## Chuẩn bị", ""] + [fill("- " + x) for x in META.get("prep", [])] + ["", "---", "", "## Nội dung từng slide", ""]
    for i, s in enumerate(SLIDES, 1):
        L.append(f"### Slide {i} · {s['part']} · {s['min']} phút")
        L.append(""); L.append("**Trình chiếu**"); L.append("")
        L.append(f"- Tiêu đề: {s['title']}")
        if s["body"]: L.append("- Nội dung: " + " · ".join(s["body"]))
        if s.get("visual"): L.append(f"- Hình/sơ đồ: dựng sẵn trong file .pptx ({s['visual'].split(':')[0]})")
        if s.get("src"): L.append(f"- Nguồn: {s['src']}")
        L += ["", "**Lời giảng**", "", s["notes"], ""]
        if s.get("example"): L += ["**Ví dụ minh họa**", "", s["example"], ""]
        if s.get("interact"):
            it = s["interact"]
            L += ["**Câu hỏi / hoạt động tương tác**", "", f"- Hình thức: {it['form']} · Thời gian: {it['time']} · Phiên bản online/lớp đông: {it['online']}",
                  f"- Câu hỏi hoặc nhiệm vụ: {it['task']}", f"- Đáp án / câu trả lời dự kiến: {it['answer']}", f"- Lời dẫn ra: {it['out']}", ""]
        L += ["---", ""]
    L += ["## Co giãn thời gian", ""] + [fill("- " + x) for x in META.get("flex", [])] + ["",
          "## Marker cần giảng viên xử lý", ""]
    marks = []
    for i, s in enumerate(SLIDES, 1):
        blob = s["title"] + " ".join(s["body"]) + s["notes"] + s.get("src", "") + str(s.get("interact", ""))
        for m in re.findall(r"\[(?:VERIFY|NEEDS PROFESSOR INPUT)[^\]]*\]", blob): marks.append(f"- Slide {i}: {m}")
    L += list(dict.fromkeys(marks))
    open(f"{OUT}/{PREFIX}_bai_giang_chi_tiet.md", "w", encoding="utf-8").write("\n".join(L) + "\n")
    print("total", total, "act", act, "talk", talk, "slides", len(SLIDES))

build_pptx(); build_md()
